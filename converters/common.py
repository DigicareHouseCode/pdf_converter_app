"""
Shared extraction machinery.

Every converter used to carry its own copy of the prompt, its own generation
config, its own JSON parser and its own DataFrame builder. They drifted apart,
so a fix applied in one place silently missed the other three. This module is
the single implementation of all of it.

Design notes worth knowing before editing:

* Values coming back from the model may be JSON numbers OR strings. Every
  coercion here goes through as_text()/as_amount(), which tolerate both. The
  previous code called .strip() directly and silently dropped the whole row on
  an AttributeError.
* Rows carry _page and _row so statement order can be restored after
  concurrent page processing. Never sort on date alone.
* A parse failure records the row in `rejected` rather than discarding it.
  Missing rows are far harder for a user to notice than wrong ones.
"""

import json
import logging
import os
import re
import time
from contextlib import contextmanager
from datetime import datetime

from config import Config

log = logging.getLogger(__name__)

AMOUNT_TOLERANCE = 0.011  # cents, with room for float noise


# ---------------------------------------------------------------- prompts ---

PROMPT_TRANSACTIONS = """
You are extracting transactions from one page of a bank statement.

STEP 1 — READ THE TABLE STRUCTURE
Locate the transaction table. Identify which visual column is which:
Date, Description, Debit (Withdrawal/Payment), Credit (Deposit/Addition),
and the running Balance. Do this before reading any values.

STEP 2 — EXTRACT EVERY TRANSACTION ROW
- Skip summary rows: "Beginning Balance", "Ending Balance", "Total", subtotals.
- Merge a description that wraps across lines into one string.
- Record the running Balance exactly as printed, when the page shows one.
- If separate Debit and Credit columns exist, put each value in its own field.
- If there is a single Amount column, use the sign, brackets, or the column
  header to decide which of Debit/Credit it belongs to.
- Never take a value from the Balance column and report it as Debit or Credit.
- Transcribe only what is visible. Do not infer or complete missing values.

STEP 3 — RETURN JSON
Return ONLY a JSON array. No markdown, no commentary.
[
  {
    "Date": "date exactly as printed on the page",
    "Description": "full description",
    "Debit": "amount or empty string",
    "Credit": "amount or empty string",
    "Balance": "running balance or empty string"
  }
]
Strip currency symbols and thousands separators from all three amounts.
"""

PROMPT_AUDIT = """
You are auditing an extraction against the page image it came from.

Check, in order:
1. COMPLETENESS  — is any transaction visible on the page missing from the
   JSON? Add it.
2. CLASSIFICATION — is any Debit sitting in Credit, or vice versa? Correct it.
3. BALANCE BLEED — was any value taken from the Balance column and reported as
   a Debit or Credit? Correct it.
4. TRANSCRIPTION — are the digits and descriptions right? Correct them.

EXTRACTED JSON:
{payload}

Return the COMPLETE corrected JSON array in the same shape. Raw JSON only.
"""


def custom_prompt(user_instructions):
    """Prompt for the user-defined extraction mode."""
    return f"""
You are extracting structured data from one page of a document, following the
user's instructions exactly.

USER INSTRUCTIONS:
{user_instructions}

RULES
- Transcribe only what is visible. Do not infer, complete, or invent values.
- Merge text that wraps across lines into one string.
- Strip currency symbols and thousands separators from numbers.
- Use the same keys on every object you return.

Return ONLY a JSON array of objects. No markdown, no commentary.
"""


# ------------------------------------------------------------------ model ---

def build_model(api_key, json_mode=True):
    """
    Build a Gemini model from config.

    Model name, temperature and token limit all come from Config so there is
    one place to change them. Three of the four converters used to hardcode
    the model name, which made Config.GEMINI_MODEL_NAME a lie.
    """
    import google.generativeai as genai

    if not api_key:
        raise ValueError("GEMINI_API_KEY is not set. Add it to your .env file.")

    genai.configure(api_key=api_key)

    generation_config = {
        "temperature": Config.GEMINI_TEMPERATURE,
        "max_output_tokens": Config.GEMINI_MAX_OUTPUT_TOKENS,
    }
    if json_mode:
        generation_config["response_mime_type"] = "application/json"

    model_name = getattr(Config, 'GEMINI_MODEL_NAME', 'gemini-2.5-pro')
    log.info("Extraction model: %s at %s DPI", model_name, Config.EXTRACTION_DPI)
    return genai.GenerativeModel(model_name, generation_config=generation_config)


class ModelUnavailable(RuntimeError):
    """The configured model name was rejected by the API."""


# finish_reason as the API reports it. 2 (MAX_TOKENS) is the one worth naming:
# it is not a transient fault and no amount of retrying clears it.
_FINISH_MAX_TOKENS = 2


def response_text(response):
    """
    Read a response's text without the `.text` quick accessor.

    That accessor raises when a candidate came back with no text part, which
    is precisely the case worth reporting rather than triggering. Returns
    (text, finish_reason).
    """
    candidates = getattr(response, 'candidates', None) or []
    if not candidates:
        return "", None
    candidate = candidates[0]
    content = getattr(candidate, 'content', None)
    parts = getattr(content, 'parts', None) or []
    text = "".join(getattr(part, 'text', '') or '' for part in parts)
    return text, getattr(candidate, 'finish_reason', None)


def generate(model, parts, retries=3, timeout=300):
    """Call the model, retrying transient failures. Returns text or ''."""
    for attempt in range(1, retries + 1):
        try:
            response = model.generate_content(parts, request_options={"timeout": timeout})
            text, finish_reason = response_text(response)
            if text:
                return text

            if finish_reason == _FINISH_MAX_TOKENS:
                # The budget went entirely on the model's own reasoning before
                # it wrote any output. Temperature is 0, so a retry buys the
                # identical result at the cost of another full call.
                log.error(
                    "Model exhausted GEMINI_MAX_OUTPUT_TOKENS (%s) before returning "
                    "any text. Raise it; retrying cannot help.",
                    Config.GEMINI_MAX_OUTPUT_TOKENS,
                )
                return ""

            log.warning("Empty model response, finish_reason=%s (attempt %s/%s)",
                        finish_reason, attempt, retries)
        except Exception as exc:
            message = str(exc)
            if "404" in message or "not found" in message.lower():
                raise ModelUnavailable(
                    f"Model '{getattr(Config, 'GEMINI_MODEL_NAME', '?')}' is not available "
                    f"for this API key."
                ) from exc
            # Quota exhaustion needs a much longer back-off than a blip.
            if "429" in message or "exhausted" in message.lower():
                time.sleep(10 * attempt)
            else:
                log.warning("Model call failed (attempt %s/%s): %s", attempt, retries, message[:200])
                time.sleep(2 * attempt)
    return ""


def render_page(page):
    """Rasterise a pdfplumber page at the configured resolution."""
    return page.to_image(resolution=Config.EXTRACTION_DPI).original


# ------------------------------------------------------------- documents ---
#
# A conversion source is either a PDF or an image. Rather than branch on the
# file type in every converter, an image is wrapped in objects that present the
# small slice of the pdfplumber interface the pipeline actually uses:
# `document.pages`, `page.extract_text()`, `page.extract_tables()` and
# `page.to_image(resolution).original`. Everything downstream — rendering,
# vision extraction, the year hint — then works unchanged.

# Derived from the upload allow-list rather than restated, so adding a
# format in one place cannot leave the converter refusing to recognise it.
IMAGE_EXTENSIONS = {"." + extension for extension in Config.ALLOWED_IMAGE_EXTENSIONS}


def is_image_file(path):
    return os.path.splitext(str(path))[1].lower() in IMAGE_EXTENSIONS


class _RenderedImage:
    """Stands in for pdfplumber's PageImage, which exposes `.original`."""

    def __init__(self, image):
        self.original = image


class ImagePage:
    """One image presented as a page."""

    def __init__(self, image, page_number):
        self._image = image
        self.page_number = page_number

    def extract_text(self):
        # An image has no text layer. Returning "" rather than raising is what
        # makes the pipeline's text-layer pass fall through to vision.
        return ""

    def extract_tables(self, table_settings=None):
        return []

    def to_image(self, resolution=None):
        # The image is already raster; `resolution` has nothing to act on.
        return _RenderedImage(self._image)


class ImageDocument:
    """A pdfplumber-shaped wrapper around one image file."""

    def __init__(self, pages):
        self.pages = pages

    def close(self):
        for page in self.pages:
            try:
                page._image.close()
            except Exception:
                pass

    def __enter__(self):
        return self

    def __exit__(self, *exc_info):
        self.close()
        return False


def load_image_pages(path):
    """
    Read an image file into one page per frame.

    Multi-frame TIFFs are common for scanned multi-page documents, and a GIF
    can carry several frames too, so every frame becomes a page. Frames are
    converted to RGB because the Gemini API rejects palette-mode images.
    """
    from PIL import Image, ImageSequence

    pages = []
    with Image.open(path) as source:
        for number, frame in enumerate(ImageSequence.Iterator(source), start=1):
            pages.append(ImagePage(frame.convert("RGB"), number))
    if not pages:
        raise ValueError("The image file contains no readable frames.")
    return pages


@contextmanager
def open_document(path):
    """
    Open a PDF or an image as a page-bearing document.

    Used as a context manager so the caller closes either kind the same way.
    """
    if is_image_file(path):
        document = ImageDocument(load_image_pages(path))
        try:
            yield document
        finally:
            document.close()
    else:
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            yield pdf


# ------------------------------------------------------------- coercion ----

def as_text(value):
    """
    Coerce any JSON scalar to a stripped string.

    Handles the case that used to destroy rows: the model returning a number
    where the schema suggested a string.
    """
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


_AMOUNT_JUNK = re.compile(r"[^\d.\-]")


def as_amount(value):
    """
    Parse a monetary value to a float, or None when there isn't one.

    Understands "$1,234.56", "(45.00)" for negatives, "1.234,56" European
    style, and plain JSON numbers.
    """
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)

    raw = str(value).strip()
    if not raw:
        return None

    negative = raw.startswith("(") and raw.endswith(")")
    if raw.lstrip("-").startswith("$"):
        raw = raw.replace("$", "")

    # European decimal comma: 1.234,56 -> 1234.56
    if re.search(r"\d,\d{2}$", raw) and "." in raw:
        raw = raw.replace(".", "").replace(",", ".")
    else:
        raw = raw.replace(",", "")

    cleaned = _AMOUNT_JUNK.sub("", raw)
    if cleaned in ("", "-", ".", "-."):
        return None

    try:
        amount = float(cleaned)
    except ValueError:
        return None

    return -abs(amount) if negative else amount


# ---------------------------------------------------------------- parsing ---

def parse_json_rows(text):
    """Pull a list of dicts out of a model response, tolerating stray prose."""
    if not text:
        return []

    cleaned = text.strip()
    if cleaned.startswith("```"):
        cleaned = re.sub(r"^```(?:json)?|```$", "", cleaned, flags=re.MULTILINE).strip()

    def _coerce(obj):
        if isinstance(obj, list):
            return [row for row in obj if isinstance(row, dict)]
        if isinstance(obj, dict):
            for key in ("transactions", "data", "rows", "items", "results"):
                if isinstance(obj.get(key), list):
                    return [row for row in obj[key] if isinstance(row, dict)]
            return [obj]
        return []

    try:
        return _coerce(json.loads(cleaned))
    except json.JSONDecodeError:
        pass

    # Salvage the outermost array, then retry after removing trailing commas.
    start, end = cleaned.find("["), cleaned.rfind("]")
    if start != -1 and end > start:
        candidate = cleaned[start:end + 1]
        for attempt in (candidate, re.sub(r",(\s*[}\]])", r"\1", candidate)):
            try:
                return _coerce(json.loads(attempt))
            except json.JSONDecodeError:
                continue

    log.warning("Could not parse model response as JSON (%d chars)", len(cleaned))
    return []


_KEYS = {
    "Date": ("Date", "date", "DATE", "TransactionDate", "transaction_date", "Posted", "PostingDate"),
    "Description": ("Description", "description", "DESC", "Details", "details",
                    "Narration", "narration", "Particulars", "Memo", "memo"),
    "Debit": ("Debit", "debit", "DEBIT", "Withdrawal", "withdrawal", "Withdrawals",
              "Payment", "payment", "Money Out", "Paid Out"),
    "Credit": ("Credit", "credit", "CREDIT", "Deposit", "deposit", "Deposits",
               "Addition", "addition", "Money In", "Paid In"),
    "Balance": ("Balance", "balance", "BALANCE", "RunningBalance", "running_balance", "Closing Balance"),
}


def _first_key(item, names):
    for name in names:
        if name in item:
            return item[name]
    return None


def normalise_rows(items, page_no, rejected=None):
    """
    Turn raw model output into canonical rows.

    Returns a list of dicts with Date, Description, Debit, Credit, Balance and
    the _page/_row ordering keys. Anything unparseable is appended to
    `rejected` rather than dropped on the floor.
    """
    rows = []

    for index, item in enumerate(items):
        if not isinstance(item, dict):
            continue

        try:
            date = as_text(_first_key(item, _KEYS["Date"]))
            description = as_text(_first_key(item, _KEYS["Description"]))
            debit = as_amount(_first_key(item, _KEYS["Debit"]))
            credit = as_amount(_first_key(item, _KEYS["Credit"]))
            balance = as_amount(_first_key(item, _KEYS["Balance"]))

            # Single-amount layouts: infer the side from sign or an explicit type.
            if debit is None and credit is None:
                amount = as_amount(_first_key(item, ("Amount", "amount", "AMOUNT", "Value", "value")))
                kind = as_text(_first_key(item, ("Type", "type", "TYPE", "TransactionType"))).upper()
                if amount is not None:
                    is_debit = (
                        amount < 0
                        or "DEBIT" in kind
                        or "WITHDRAW" in kind
                        or "PAYMENT" in kind
                    )
                    if is_debit:
                        debit = abs(amount)
                    else:
                        credit = abs(amount)

            # Amounts are magnitudes; the column carries the direction.
            if debit is not None:
                debit = abs(debit)
            if credit is not None:
                credit = abs(credit)

            if not date and not description:
                continue

            rows.append({
                "Date": date,
                "Description": re.sub(r"\s+", " ", description).strip(),
                "Debit": debit,
                "Credit": credit,
                "Balance": balance,
                "_page": page_no,
                "_row": index,
                "_flags": [],
            })
        except Exception as exc:  # never lose a row silently
            log.warning("Row %s on page %s could not be normalised: %s", index, page_no, exc)
            if rejected is not None:
                rejected.append({"page": page_no, "row": index, "raw": item, "error": str(exc)})

    return rows


# ------------------------------------------------------------------ dates ---

_YEAR = re.compile(r"\b(19|20)\d{2}\b")
_DATE_FORMATS = (
    "%m/%d/%Y", "%d/%m/%Y", "%Y-%m-%d", "%m-%d-%Y", "%d-%m-%Y",
    "%d %b %Y", "%b %d %Y", "%d %B %Y", "%B %d %Y",
    "%m/%d/%y", "%d/%m/%y",
)
_DATE_FORMATS_NO_YEAR = ("%m/%d", "%d/%m", "%b %d", "%d %b", "%m-%d")

# The same formats, preferring day-first. Which list is used is decided per
# document by _infer_day_first(), because "12-01-2026" is a valid date under
# either reading and guessing wrong silently moves a transaction by months.
_DATE_FORMATS_DAY_FIRST = (
    "%d/%m/%Y", "%m/%d/%Y", "%Y-%m-%d", "%d-%m-%Y", "%m-%d-%Y",
    "%d %b %Y", "%b %d %Y", "%d %B %Y", "%B %d %Y",
    "%d/%m/%y", "%m/%d/%y",
)
_DATE_FORMATS_NO_YEAR_DAY_FIRST = ("%d/%m", "%m/%d", "%d %b", "%b %d", "%d-%m")

_NUMERIC_DATE = re.compile(r"^\s*(\d{1,4})[/\-.](\d{1,2})[/\-.](\d{2,4})\s*$")


def _infer_day_first(rows):
    """
    Decide whether this document writes numeric dates day-first or month-first.

    One date on its own is ambiguous, but a statement rarely is: a single row
    whose first component exceeds 12 settles the reading for every other row.
    Ambiguous documents keep the month-first default.
    """
    day_first = month_first = 0

    for row in rows:
        match = _NUMERIC_DATE.match(row.get("Date") or "")
        if not match or len(match.group(1)) == 4:  # 4 digits means Y-M-D already
            continue
        first, second = int(match.group(1)), int(match.group(2))
        if first > 12 >= second:
            day_first += 1
        elif second > 12 >= first:
            month_first += 1

    if day_first or month_first:
        log.info("Date order: %s (%d day-first vs %d month-first signals)",
                 "day-first" if day_first > month_first else "month-first",
                 day_first, month_first)
    return day_first > month_first


def statement_year(pdf):
    """Best guess at the statement's year, read from the first two pages."""
    try:
        for page in pdf.pages[:2]:
            text = page.extract_text() or ""
            years = [int(m.group(0)) for m in _YEAR.finditer(text)]
            plausible = [y for y in years if 1990 <= y <= datetime.now().year + 1]
            if plausible:
                return max(plausible)
    except Exception:
        pass
    return None


def resolve_dates(rows, year_hint=None):
    """
    Parse dates and fill in a missing year.

    Replaces the old behaviour of appending a hardcoded '/2025' to every date,
    which mis-sorted any statement from another year and produced NaT for any
    date that already carried one.

    Where the year is absent, it is taken from `year_hint`. A month that goes
    backwards relative to its predecessor is treated as a year rollover.

    Numeric dates are read day-first or month-first according to the document
    as a whole — see _infer_day_first().
    """
    fallback = year_hint or datetime.now().year
    previous_month = None
    year = fallback

    if _infer_day_first(rows):
        formats = _DATE_FORMATS_DAY_FIRST
        formats_no_year = _DATE_FORMATS_NO_YEAR_DAY_FIRST
    else:
        formats = _DATE_FORMATS
        formats_no_year = _DATE_FORMATS_NO_YEAR

    for row in rows:
        raw = (row.get("Date") or "").strip()
        row["_date"] = None
        if not raw:
            continue

        parsed = None
        for fmt in formats:
            try:
                parsed = datetime.strptime(raw, fmt)
                break
            except ValueError:
                continue

        if parsed is None:
            for fmt in formats_no_year:
                try:
                    partial = datetime.strptime(raw, fmt)
                except ValueError:
                    continue
                # December -> January means the statement crossed a year end.
                if previous_month is not None and partial.month < previous_month - 6:
                    year += 1
                parsed = partial.replace(year=year)
                break

        if parsed is not None:
            previous_month = parsed.month
            row["_date"] = parsed
            row["Date"] = parsed.strftime("%Y-%m-%d")

    return rows


def order_rows(rows):
    """Restore statement order after concurrent page processing."""
    return sorted(rows, key=lambda r: (r.get("_page", 0), r.get("_row", 0)))


# ---------------------------------------------------------- reconciliation ---

def reconcile(rows):
    """
    Verify the running balance across consecutive rows.

    For each pair where both balances are known:

        balance[n] == balance[n-1] - debit[n] + credit[n]

    This is the strongest signal a bank statement offers, and it catches the
    four failure modes that matter: swapped debit/credit, a missing row, a
    misread digit, and a balance leaking into an amount column.

    Returns (checked, flagged).
    """
    checked = 0
    flagged = 0
    previous = None

    for row in rows:
        balance = row.get("Balance")
        if balance is None:
            previous = None if balance is None else balance
            continue

        if previous is not None:
            debit = row.get("Debit") or 0.0
            credit = row.get("Credit") or 0.0
            expected = previous - debit + credit
            checked += 1
            drift = balance - expected
            if abs(drift) > AMOUNT_TOLERANCE:
                flagged += 1
                row["_flags"].append(
                    f"Balance off by {drift:+,.2f} (expected {expected:,.2f})"
                )

        previous = balance

    return checked, flagged


def summarise(rows, checked, flagged, rejected_count=0):
    """Totals for the summary sheet and the result page."""
    debits = sum(r["Debit"] for r in rows if r.get("Debit"))
    credits = sum(r["Credit"] for r in rows if r.get("Credit"))
    return {
        "rows": len(rows),
        "debits": debits,
        "credits": credits,
        "net": credits - debits,
        "checked": checked,
        "flagged": flagged,
        "rejected": rejected_count,
    }


# ------------------------------------------------------------------ excel ---

COLUMNS = ["Date", "Description", "Debit", "Credit", "Balance"]


def write_workbook(rows, path, summary, sheet_name="Transactions"):
    """
    Write the transactions sheet plus a summary sheet.

    Rows that failed reconciliation get a Review note and a highlighted band,
    so a mismatch is visible at a glance rather than buried in the numbers.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="0F5F56")
    flag_fill = PatternFill("solid", fgColor="FBECEC")
    money = '#,##0.00'

    headers = COLUMNS + ["Review"]
    sheet.append(headers)
    for column in range(1, len(headers) + 1):
        cell = sheet.cell(row=1, column=column)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row in rows:
        note = "; ".join(row.get("_flags", []))
        sheet.append([
            row.get("Date", ""),
            row.get("Description", ""),
            row.get("Debit"),
            row.get("Credit"),
            row.get("Balance"),
            note,
        ])
        index = sheet.max_row
        for column in (3, 4, 5):
            sheet.cell(row=index, column=column).number_format = money
        if note:
            for column in range(1, len(headers) + 1):
                sheet.cell(row=index, column=column).fill = flag_fill

    widths = [13, 52, 13, 13, 14, 40]
    for position, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(position)].width = width
    sheet.freeze_panes = "A2"

    report = workbook.create_sheet("Summary")
    report.append(["Metric", "Value"])
    report["A1"].font = header_font
    report["B1"].font = header_font
    report["A1"].fill = header_fill
    report["B1"].fill = header_fill

    for label, value, fmt in [
        ("Transactions extracted", summary["rows"], None),
        ("Total debits", summary["debits"], money),
        ("Total credits", summary["credits"], money),
        ("Net movement", summary["net"], money),
        ("Balance checks performed", summary["checked"], None),
        ("Balance mismatches flagged", summary["flagged"], None),
        ("Rows that could not be read", summary["rejected"], None),
    ]:
        report.append([label, value])
        if fmt:
            report.cell(row=report.max_row, column=2).number_format = fmt

    report.column_dimensions["A"].width = 30
    report.column_dimensions["B"].width = 18

    if summary["flagged"]:
        report.append([])
        report.append([
            "Check the highlighted rows on the Transactions sheet — their running "
            "balance does not match the previous row."
        ])

    workbook.save(path)


def result_message(summary):
    """One-line outcome for the UI."""
    parts = [f"{summary['rows']} transactions extracted"]
    if summary["flagged"]:
        parts.append(f"{summary['flagged']} flagged for review")
    elif summary["checked"]:
        parts.append("running balance reconciles")
    if summary["rejected"]:
        parts.append(f"{summary['rejected']} rows unreadable")
    return ", ".join(parts) + "."
